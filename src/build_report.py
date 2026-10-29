
"""Build a results-filled DOCX report after src/train.py has been run."""
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
FIG = OUT / "figures"
report = Document()
report.styles["Normal"].font.name = "Times New Roman"
report.styles["Normal"].font.size = Pt(10.5)

p = report.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("CONCRETE COMPRESSIVE STRENGTH PREDICTION")
r.bold = True; r.font.size = Pt(18)
p = report.add_paragraph("Comparative Performance Analysis of Machine Learning Models")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
report.add_paragraph("Student: ____________________    Roll No.: ____________________")
report.add_paragraph("Course / Faculty: ____________________    Date: ____________________")

def h(text, level=1): report.add_heading(text, level=level)
def ptext(text): report.add_paragraph(text)
def add_image(filename, caption, width=6.0):
    path = FIG / filename
    if path.exists():
        report.add_picture(str(path), width=Inches(width))
        cp = report.add_paragraph(caption)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

summary_path = OUT / "results_summary.md"
if not summary_path.exists():
    raise SystemExit("Run `python src/train.py` first. Missing outputs/results_summary.md.")
audit = __import__("json").loads((OUT/"data_audit.json").read_text())
baseline = pd.read_csv(OUT/"baseline_model_comparison.csv")
final = pd.read_csv(OUT/"final_model_comparison.csv")
tuning = pd.read_csv(OUT/"tuning_comparison.csv")
details = __import__("json").loads((OUT/"tuning_details.json").read_text())

h("Abstract")
ptext(f"This project predicts concrete compressive strength in MPa using the UCI Concrete Compressive Strength dataset. The executed dataset contains {audit['rows']} observations and {len(audit['dtypes'])-1} predictors. Linear Regression, Random Forest and Gradient Boosting were compared using a common held-out test set; Random Forest hyperparameters were tuned using training-only cross-validation. The final selected model was {final.iloc[0]['Model']}, with RMSE {final.iloc[0]['RMSE (MPa)']:.3f} MPa, MAE {final.iloc[0]['MAE (MPa)']:.3f} MPa and R² {final.iloc[0]['R²']:.3f}. Results are specific to the documented split and should not be interpreted as engineering certification.")
h("1. Introduction and Problem Statement")
ptext("Concrete compressive strength is influenced by mixture proportions and curing age. The project estimates strength from seven material quantities and age. The target is continuous, so this is a supervised regression problem.")
h("2. Objectives")
for item in ["Perform EDA and data-quality checks.","Compare three regression models fairly.","Use leakage-aware preprocessing and training-only cross-validation.","Tune a model and evaluate whether performance improves.","Create report-ready charts, results and a Streamlit prediction application."]:
    report.add_paragraph(item, style="List Bullet")
h("3. Literature Review")
ptext("Yeh (1998) introduced the dataset and investigated neural-network modelling of high-performance concrete strength. Linear regression provides an interpretable baseline. Random Forest combines randomized decision trees (Breiman, 2001), while Gradient Boosting builds an additive sequence of learners to reduce residual error (Friedman, 2001). Scikit-learn pipelines and cross-validation support reproducible model evaluation (Pedregosa et al., 2011).")
h("4. Dataset Description")
ptext("Source: UCI Machine Learning Repository, Concrete Compressive Strength, dataset ID 165. URL: https://archive.ics.uci.edu/dataset/165/concrete+compressive+strength. DOI: https://doi.org/10.24432/C5PK67. The dataset has 1,030 observations, eight quantitative inputs and one quantitative target, and is licensed CC BY 4.0.")
ptext("Inputs: Cement, Blast Furnace Slag, Fly Ash, Water, Superplasticizer, Coarse Aggregate and Fine Aggregate (kg/m³), and Age (days). Target: Concrete compressive strength (MPa).")
h("5. EDA and Data Quality")
ptext(f"The executed data audit found {audit['rows']} rows, {audit['columns']} columns, {audit['missing_cells']} missing cells after numeric conversion, and {audit['exact_duplicate_rows']} exact duplicate rows. Duplicate records were reported rather than automatically deleted because exact equality alone does not prove an observation is erroneous.")
add_image("target_distribution.png", "Figure 1. Target distribution.")
add_image("feature_distributions.png", "Figure 2. Input feature distributions.")
add_image("correlation_heatmap.png", "Figure 3. Pearson correlation heatmap.")
add_image("feature_boxplots.png", "Figure 4. Input feature boxplots. Features have different scales; use for descriptive inspection.")
add_image("feature_vs_target.png", "Figure 5. Input features versus target.")
h("6. Preprocessing and Leakage Prevention")
ptext("An 80/20 train/test split with random_state=42 is used. Median imputation and scaling are pipeline steps, so learned preprocessing is fit on training data only. Linear Regression is scaled; tree-based models are not scaled. RandomizedSearchCV uses five-fold cross-validation only within the training subset. The held-out test set is not used to select hyperparameters.")
h("7. Models and Experimental Setup")
ptext("The baseline models are Linear Regression, Random Forest Regressor and Gradient Boosting Regressor. All share the same split and are evaluated using MAE, MSE, RMSE and R². Classification metrics such as accuracy, precision, recall, F1 and ROC-AUC are not suitable for the continuous target in this regression formulation.")
h("8. Baseline Performance")
t = report.add_table(rows=1, cols=len(baseline.columns)); t.style = "Light Shading Accent 1"
for i,c in enumerate(baseline.columns): t.rows[0].cells[i].text = str(c)
for _, row in baseline.iterrows():
    cells=t.add_row().cells
    for i,c in enumerate(baseline.columns):
        v=row[c]
        cells[i].text = f"{v:.4f}" if isinstance(v,(int,float)) else str(v)
add_image("comparison_mae.png", "Figure 6. Baseline/final model MAE comparison.")
add_image("comparison_rmse.png", "Figure 7. Baseline/final model RMSE comparison.")
add_image("comparison_r2.png", "Figure 8. Baseline/final model R² comparison.")
h("9. Hyperparameter Tuning")
ptext(f"RandomizedSearchCV searched 25 sampled Random Forest configurations using five-fold CV and negative RMSE scoring. Best training CV RMSE: {-details['best_cv_RMSE_MPa']:.3f} MPa. Selected parameters: {details['best_params']}.")
t = report.add_table(rows=1, cols=len(tuning.columns)); t.style = "Light Shading Accent 1"
for i,c in enumerate(tuning.columns): t.rows[0].cells[i].text = str(c)
for _, row in tuning.iterrows():
    cells=t.add_row().cells
    for i,c in enumerate(tuning.columns):
        v=row[c]
        cells[i].text = f"{v:.4f}" if isinstance(v,(int,float)) else str(v)
h("10. Final Comparative Analysis")
t = report.add_table(rows=1, cols=len(final.columns)); t.style = "Light Shading Accent 1"
for i,c in enumerate(final.columns): t.rows[0].cells[i].text = str(c)
for _, row in final.iterrows():
    cells=t.add_row().cells
    for i,c in enumerate(final.columns):
        v=row[c]
        cells[i].text = f"{v:.4f}" if isinstance(v,(int,float)) else str(v)
winner=final.iloc[0]
ptext(f"The lowest-RMSE candidate on this fixed holdout split is {winner['Model']}. It achieved MAE {winner['MAE (MPa)']:.3f} MPa, MSE {winner['MSE (MPa²)']:.3f} MPa², RMSE {winner['RMSE (MPa)']:.3f} MPa and R² {winner['R²']:.3f}. RMSE was used as the primary ranking metric, while MAE and R² were considered together to avoid relying on a single metric.")
add_image("actual_vs_predicted.png", "Figure 9. Actual versus predicted strength for the selected model.")
add_image("residual_plot.png", "Figure 10. Residuals versus predicted strength.")
add_image("permutation_importance.png", "Figure 11. Permutation importance. This is predictive importance, not causal evidence.")
h("11. Streamlit Deployment")
ptext("The app is implemented in app.py. Run `streamlit run app.py` after training. It loads the saved model, accepts eight numerical inputs and displays a point estimate in MPa. Add the hosted URL after deployment: ____________________. The app warns that predictions are educational estimates and not a replacement for laboratory testing or engineering approval.")
h("12. Results and Discussion")
ptext("The selected model and its scores are based on one fixed random holdout split. Results can vary under different splits or dataset versions. A random row split can also place related mix designs at different ages in both train and test subsets, so scores may not reflect performance on entirely unseen mix designs. Feature importance should not be interpreted as causation.")
h("13. Conclusion and Future Scope")
ptext("The project demonstrates a reproducible workflow for concrete strength regression, including EDA, leakage-aware preprocessing, baseline comparison, tuning, plots and a deployment interface. Future work should evaluate grouped splits by mixture composition, repeated cross-validation, uncertainty estimates and independent laboratory data. Do not use this model for structural design or compliance decisions without independent validation, qualified engineering review and required testing.")
h("14. GitHub Repository")
ptext("Repository URL after upload: ____________________. Include source code, notebook, requirements, report, dataset source link, README and generated model comparison results/figures.")
h("15. References")
for ref in [
"Yeh, I.-C. (1998). Modeling of strength of high-performance concrete using artificial neural networks. Cement and Concrete Research, 28(12), 1797–1808. https://doi.org/10.1016/S0008-8846(98)00165-3",
"Yeh, I.-C. (1998). Concrete Compressive Strength [Dataset]. UCI Machine Learning Repository. https://doi.org/10.24432/C5PK67",
"Breiman, L. (2001). Random Forests. Machine Learning, 45, 5–32. https://doi.org/10.1023/A:1010933404324",
"Friedman, J. H. (2001). Greedy Function Approximation: A Gradient Boosting Machine. The Annals of Statistics, 29(5), 1189–1232. https://doi.org/10.1214/aos/1013203451",
"Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825–2830. https://jmlr.org/papers/v12/pedregosa11a.html",
"Streamlit documentation. https://docs.streamlit.io/"
]:
    report.add_paragraph(ref, style="List Number")
path = ROOT / "report" / "Project_Report_Completed.docx"
report.save(path)
print(f"Created {path}")
